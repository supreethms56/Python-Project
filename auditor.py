"""
auditor.py: Object-Oriented classes for Transaction and ReportAuditor.
Fulfills:
- AC-3: Read Transaction ID, Product Name, Category, Price, and Quantity.
- AC-5: Create Total Amount column when it does not exist.
- AC-7: Create Discounted Total column for final values.
- AC-8: Continue processing valid records when invalid records exist.
- AC-9: Flag and log rows with missing or invalid data.
- AC-10: Do not terminate unexpectedly because of invalid data.
- AC-11: Implement object-oriented design using classes.
- AC-12: Modular architecture.
- AC-13: Generate a bar chart (Product Name vs Discounted Total).
- AC-14: Save processed workbook as audit_<original_filename>.xlsx.
"""

from pathlib import Path
from collections import defaultdict
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

from utils import calculate_transaction_totals


class Transaction:
    """AC-11: Encapsulates individual transaction row data."""

    def __init__(
        self,
        transaction_id: Any,
        product_name: str,
        category: str,
        price: float,
        quantity: float,
        total_amount: float,
        discounted_total: float,
    ):
        self.transaction_id = transaction_id
        self.product_name = str(product_name).strip() if product_name else "Unknown"
        self.category = str(category).strip() if category else "Uncategorized"
        self.price = price
        self.quantity = quantity
        self.total_amount = total_amount
        self.discounted_total = discounted_total


class ReportAuditor:
    """AC-11: Encapsulates spreadsheet parsing, validation, enrichment, and chart rendering."""

    REQUIRED_HEADERS = ["Transaction_ID", "Product_Name", "Category", "Price", "Quantity"]

    def __init__(self, file_path: Path):
        self.file_path = Path(file_path)
        self.workbook = None
        self.sheet = None
        self.valid_transactions = []
        self.error_logs = []
        self.file_revenue = 0.0
        self.product_units_sold = defaultdict(float)

    def process_file(self) -> bool:
        """
        Executes audit flow for a single workbook.
        Returns True if the file was processed, False if it was critically unreadable.
        """
        # AC-8 & AC-10: Graceful file-level error handling
        try:
            self.workbook = openpyxl.load_workbook(self.file_path)
            self.sheet = self.workbook.active
        except Exception as exc:
            self._record_error(f"Failed to open workbook: {exc}")
            self._write_error_log()
            return False

        header_indices = self._get_header_indices()

        # AC-3: Ensure required columns exist
        missing_columns = [col for col in self.REQUIRED_HEADERS if col not in header_indices]
        if missing_columns:
            self._record_error(f"Missing required columns {missing_columns}. Processing aborted for this file.")
            self._write_error_log()
            return False

        # AC-5 & AC-7: Locate or create 'Total Amount' and 'Discounted Total' columns
        total_amt_col = self._get_or_create_column("Total Amount", header_indices)
        disc_total_col = self._get_or_create_column("Discounted Total", header_indices)

        # AC-3, AC-8, AC-9, AC-10: Iterate rows, process valid data, log invalid records
        for row_idx in range(2, self.sheet.max_row + 1):
            trans_id = self.sheet.cell(row=row_idx, column=header_indices["Transaction_ID"]).value
            prod_name = self.sheet.cell(row=row_idx, column=header_indices["Product_Name"]).value
            category = self.sheet.cell(row=row_idx, column=header_indices["Category"]).value
            price = self.sheet.cell(row=row_idx, column=header_indices["Price"]).value
            quantity = self.sheet.cell(row=row_idx, column=header_indices["Quantity"]).value

            total_amount, disc_total, error_msg = calculate_transaction_totals(price, quantity)

            if error_msg:
                # AC-9 & AC-10: Flag and log invalid row, continue execution
                self._record_error(f"Row {row_idx} [ID: {trans_id}]: {error_msg}")
                continue

            # Populate computed values into the sheet
            # self.sheet.cell(row=row_idx, column=total_amt_col, value=total_amount)
            # self.sheet.cell(row=row_idx, column=disc_total_col, value=disc_total)


            # Populate computed values into the sheet and apply currency formatting
            cell_total = self.sheet.cell(row=row_idx, column=total_amt_col, value=total_amount)
            cell_disc = self.sheet.cell(row=row_idx, column=disc_total_col, value=disc_total)

            cell_total.number_format = '"$"#,##0.00'
            cell_disc.number_format = '"$"#,##0.00'

            # Record valid transaction
            tx = Transaction(trans_id, prod_name, category, float(price), float(quantity), total_amount, disc_total)
            self.valid_transactions.append(tx)
            self.file_revenue += disc_total
            self.product_units_sold[tx.product_name] += tx.quantity

        # AC-13: Generate Bar Chart in workbook
        self._add_bar_chart(header_indices["Product_Name"], disc_total_col)

        # AC-14: Save output file as audit_<original_filename>.xlsx
        output_filename = f"audit_{self.file_path.name}"
        output_path = self.file_path.parent / output_filename
        try:
            self.workbook.save(output_path)
            print(f"  -> Saved audited workbook: {output_path.name}")
        except Exception as exc:
            self._record_error(f"Failed to save {output_filename}: {exc}")

        # Write error log file if anomalies were detected
        self._write_error_log()
        return True

    def _get_header_indices(self) -> dict:
        """Reads row 1 and maps trimmed column header titles to their column index."""
        headers = {}
        for col_idx in range(1, self.sheet.max_column + 1):
            val = self.sheet.cell(row=1, column=col_idx).value
            if val is not None:
                headers[str(val).strip()] = col_idx
        return headers

    def _get_or_create_column(self, col_title: str, header_indices: dict) -> int:
        """Finds existing column by title or appends it to row 1 (AC-5, AC-7)."""
        if col_title in header_indices:
            return header_indices[col_title]

        new_col_idx = self.sheet.max_column + 1
        self.sheet.cell(row=1, column=new_col_idx, value=col_title)
        header_indices[col_title] = new_col_idx
        return new_col_idx

    def _add_bar_chart(self, prod_name_col: int, disc_total_col: int):
        """Creates and embeds a Bar Chart plotting Product Name vs Discounted Total (AC-13)."""
        if self.sheet.max_row <= 1 or not self.valid_transactions:
            return

        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Product Name vs. Discounted Total"
        chart.y_axis.title = "Discounted Total ($)"
        chart.x_axis.title = "Product Name"

        # Reference data and category headers
        data = Reference(self.sheet, min_col=disc_total_col, min_row=1, max_row=self.sheet.max_row)
        categories = Reference(self.sheet, min_col=prod_name_col, min_row=2, max_row=self.sheet.max_row)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.legend = None

        # Position chart next to data table
        target_cell = f"{get_column_letter(self.sheet.max_column + 2)}2"
        self.sheet.add_chart(chart, target_cell)

    def _record_error(self, message: str):
        self.error_logs.append(message)

    def _write_error_log(self):
        """AC-9: Records corrupt records in a separate log text file."""
        if not self.error_logs:
            return
        log_path = self.file_path.parent / f"error_log_{self.file_path.stem}.txt"
        with open(log_path, "w", encoding="utf-8") as f:
            for entry in self.error_logs:
                f.write(f"[{self.file_path.name}] {entry}\n")